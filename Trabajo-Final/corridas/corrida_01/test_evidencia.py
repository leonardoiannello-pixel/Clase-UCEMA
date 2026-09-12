"""Read-only evidence tests: never generate, consolidate or edit review files."""
import hashlib
import json
from pathlib import Path
import unittest
from zipfile import ZipFile
import openpyxl

ROOT=Path(__file__).resolve().parent
PACKAGE=ROOT.parents[1]


def read(name):
    return json.loads((ROOT/name).read_text(encoding='utf-8'))


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


class EvidenceTests(unittest.TestCase):
    def test_effective_prompts_exact(self):
        record=read('registro.json')
        for key in ['system_prompt','user_prompt']:
            info=record[key]
            self.assertEqual(sha(ROOT/info['archivo_efectivamente_usado']),info['sha256'])
            self.assertEqual((ROOT/info['archivo_efectivamente_usado']).read_bytes(),(PACKAGE/'prompts'/(key+'.md')).read_bytes())

    def test_inputs_exact_synthetic_catalog(self):
        catalog=read('fuentes/inputs_referencias.json')['files']
        for entry in read('registro.json')['inputs']:
            self.assertEqual(sha(ROOT/entry['archivo_evidencia']),entry['sha256'])
            self.assertEqual(entry['sha256'],catalog[entry['tipo']]['sha256'])

    def test_real_cli_report_and_exit(self):
        invocation=read('cli_invocacion.json')
        self.assertEqual(invocation['exit_code'],0)
        self.assertIn('generate',invocation['argv'])
        self.assertNotIn('consolidate',invocation['argv'])
        self.assertEqual(read('cli_stdout.txt'),read('ejecucion/reporte.json'))
        self.assertEqual((ROOT/'cli_stderr.txt').read_bytes(),b'')
        self.assertEqual((ROOT/'ejecucion/tool_stderr.txt').read_bytes(),b'')

    def test_output_hashes_and_no_consolidation(self):
        record=read('registro.json')
        self.assertEqual(record['operacion'],'generate')
        self.assertFalse(record['consolidacion_ejecutada'])
        self.assertEqual(record['estado'],read('ejecucion/reporte.json')['estado'])
        for entry in record['archivos_generados']:
            self.assertEqual(sha(ROOT/entry['archivo_evidencia']),entry['sha256'])
        self.assertEqual(len(list((ROOT/'ejecucion/artifacts').glob('*.xlsx'))),5)
        self.assertFalse(list(ROOT.rglob('consolidated_final.xlsx')))

    def test_only_adjustment_editable_and_zero(self):
        manifest=read('ejecucion/artifacts/manifest.json')
        for name,info in manifest['files'].items():
            workbook=openpyxl.load_workbook(ROOT/'ejecucion/artifacts'/name)
            for sheet in workbook:
                self.assertTrue(sheet.protection.sheet)
                unlocked={c.coordinate for row in sheet for c in row if not c.protection.locked}
                expected={'R'+str(i) for i in range(2,len(info['employee_ids'])+2)} if sheet.title=='Detail' else set()
                self.assertEqual(unlocked,expected)
                for coordinate in unlocked:
                    self.assertEqual(sheet[coordinate].value,0)
            workbook.close()

    def test_budgets_and_exceptions_from_outputs(self):
        record=read('registro.json')
        report=read('ejecucion/reporte.json')
        self.assertEqual(record['budgets'],report['budgets'])
        self.assertEqual(record['excepciones'],read('ejecucion/artifacts/generation_validation.json')['events'])
        manifest=read('ejecucion/artifacts/manifest.json')
        for entry in record['archivos_revision']:
            workbook=openpyxl.load_workbook(ROOT/'ejecucion/artifacts'/entry['archivo'],data_only=True)
            values=dict(workbook['Summary'].values)
            budget=next(b for b in record['budgets'] if b['pool']==entry['pool'])
            self.assertEqual(budget['proposed_payroll'],values['Proposed Payroll'])
            self.assertEqual(budget['maximum_payroll'],values['Maximum Payroll / Budget'])
            self.assertEqual(budget['budget_status'],values['Budget Status'])
            workbook.close()

    def test_approval_and_unexposed_metrics(self):
        record=read('registro.json')
        self.assertEqual(record['estado_aprobacion'],'PENDING_HUMAN_APPROVAL')
        self.assertEqual(record['intervenciones_humanas'],[])
        self.assertFalse(record['ajustes_humanos_realizados'])
        for key in ['modelo','costo_por_corrida']:
            self.assertIsNone(record[key])
        for key in ['input_total','input_cached','output']:
            self.assertIsNone(record['tokens'][key])
        self.assertFalse(read('host_metadata.json')['llm_api_adicional_invocada'])

    def test_proposal_matches_validated_v4_values(self):
        current=read('ejecucion/artifacts/manifest.json')['proposals']
        prior=json.loads((PACKAGE.parent/'Entrega-2/v4/demo/generated/manifest.json').read_text(encoding='utf-8'))['proposals']
        self.assertEqual(current,prior)

    def test_formula_cache_no_errors(self):
        for path in (ROOT/'ejecucion/artifacts').glob('*.xlsx'):
            workbook=openpyxl.load_workbook(path,data_only=True)
            for sheet in workbook:
                for row in sheet:
                    for cell in row:
                        self.assertNotEqual(cell.data_type,'e',(path.name,sheet.title,cell.coordinate))
            workbook.close()

    def test_password_capture_control(self):
        control=read('control_password.json')
        self.assertFalse(control['contrasena_en_claro_detectada'])
        self.assertNotIn('--password',read('cli_invocacion.json')['argv'])
        self.assertNotIn('--password',read('ejecucion/tool_invocation.json')['command'])
        for path in (ROOT/'ejecucion/artifacts').glob('*.xlsx'):
            with ZipFile(path) as z:
                self.assertFalse(any('externalLinks' in name for name in z.namelist()))


if __name__=='__main__':
    unittest.main()
