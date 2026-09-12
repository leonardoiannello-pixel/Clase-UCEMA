import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile, ZIP_DEFLATED
import xml.etree.ElementTree as ET

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from engine import D, salary, weight, human_result, at_x
from workflow import ROOT, load_inputs, calculate_group, generate, consolidate, rows, digest


def change_cell(path, sheet, coordinate, value):
    """Simulate a hostile ZIP edit or a user's numeric adjustment; no Excel writer."""
    with ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    name = f'xl/worksheets/sheet{sheet}.xml'
    root = ET.fromstring(parts[name])
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    cell = next(c for c in root.iter(ns+'c') if c.get('r') == coordinate)
    for child in list(cell):
        cell.remove(child)
    cell.set('t', 'n')
    ET.SubElement(cell, ns+'v').text = str(value)
    parts[name] = ET.tostring(root, encoding='utf-8', xml_declaration=True)
    with ZipFile(path, 'w', ZIP_DEFLATED) as z:
        for n, data in parts.items():
            z.writestr(n, data)


class WorkflowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.base = Path(cls.temp.name)
        cls.original = cls.base/'original'
        cls.manifest = generate(cls.original, ROOT/'inputs/Parameters.xlsx', 'DEMO-only')
        cls.records = [e for es in cls.manifest['proposals'].values() for e in es]
        cls.by_id = {e['Employee_ID']: e for e in cls.records}

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_v3_regression_all_fields(self):
        for f in (ROOT.parent/'runs/run_3_v3/outputs').glob('*.xlsx'):
            for original in rows(f, 'Detail'):
                current = self.by_id[original['Employee_ID']]
                for field, value in original.items():
                    if field == 'Compa_Ratio_Improvement' and value is None:
                        self.assertIsNone(current['Final_Compa_Ratio'])
                    elif field == 'Compa_Ratio_Improvement':
                        self.assertAlmostEqual(float(current['Final_Compa_Ratio']-current['Initial_Compa_Ratio']), value, places=10)
                    elif field == 'Flag':
                        self.assertEqual(current[field], value or '')
                    elif value is None:
                        self.assertIsNone(current[field])
                    elif isinstance(value, (float, int)):
                        self.assertAlmostEqual(float(current[field]), value, places=7, msg=f'{original["Employee_ID"]} {field}')
                    else:
                        self.assertEqual(current[field], value)

    def test_leaders_excluded_and_mapping(self):
        import openpyxl
        for filename, info in self.manifest['files'].items():
            w = openpyxl.load_workbook(self.original/filename, data_only=True)
            ids = [r[0] for r in list(w['Detail'].values)[1:]]
            leader_ids = {e['Employee_ID'] for e in self.records if e['Is_Leader']=='YES'}
            if info['group']=='Leadership':
                self.assertEqual(set(ids), leader_ids)
                self.assertEqual(w['Summary']['B3'].value, 'CEO001')
            else:
                self.assertFalse(set(ids) & leader_ids)
                self.assertNotIn(w['Summary']['B3'].value, ids)
            w.close()

    def test_only_adjustment_editable(self):
        import openpyxl
        for filename, info in self.manifest['files'].items():
            w = openpyxl.load_workbook(self.original/filename)
            self.assertTrue(w.security.lockStructure)
            for s in w:
                self.assertTrue(s.protection.sheet)
                unlocked = {c.coordinate for row in s for c in row if not c.protection.locked}
                expected = {'R'+str(i) for i in range(2, len(info['employee_ids'])+2)} if s.title=='Detail' else set()
                self.assertEqual(unlocked, expected)
            self.assertEqual(len(w['Detail'].data_validations.dataValidation), 1)
            self.assertEqual(w.calculation.calcMode, 'auto')
            w.close()

    def test_salary_rounding(self):
        for raw, expected in [('100049.99',100000),('100050',100100),('100150',100200)]:
            self.assertEqual(salary(raw), expected)
        self.assertTrue(all(e['New_Salary'] % 100 == 0 for e in self.records))

    def test_negative_protected_floor(self):
        with self.assertRaisesRegex(ValueError, 'PROTECTED_FLOOR'):
            human_result(self.by_id['C001'], -.01, -.05, .05)
        result = human_result(self.by_id['A001'], -.01, -.05, .05)
        self.assertLess(result['Final_Salary'], self.by_id['A001']['New_Salary'])

    def test_limits_precision_and_nonfinite(self):
        for adj, error in [(.06,'OUT_OF_PARAMETERS'),(.00001,'PRECISION'),('=1+1','NOT_NUMERIC'),(True,'NOT_NUMERIC'),(float('nan'),'NOT_FINITE')]:
            with self.subTest(adj=adj), self.assertRaisesRegex(ValueError, error):
                human_result(self.by_id['A001'], adj, -.05, .05)

    def test_market_bands_and_cap(self):
        for ratio, expected in [('.69999',4),('.7',2),('.8',1),('.9',D('.5')),('1',0)]:
            self.assertEqual(weight(D(ratio)), expected)
        for e in self.records:
            if e['Market_Increase_Pct']:
                raw = D(e['June_Base_Salary'])*(1+e['Protected_Increase_Pct']+e['Merit_Increase_Pct']+e['Market_Increase_Pct'])
                self.assertLessEqual(raw, e['WTW_Reference_Salary'])
        self.assertEqual(self.by_id['B004']['Market_Increase_Pct'],0)

    def test_gamma_exception_preserved(self):
        es = self.manifest['proposals']['Team Gamma']
        self.assertEqual(sum(e['New_Salary'] for e in es), 9322000)
        self.assertTrue(all(e['X']==0 and 'BUDGET_INSUFICIENTE' in e['Flag'] for e in es))

    def reviewed_copy(self, name):
        path = self.base/name
        shutil.copytree(self.original, path)
        return path

    def test_consolidation_preserves_original_and_budgets(self):
        reviewed = self.reviewed_copy('valid')
        change_cell(reviewed/'team_L-A.xlsx',2,'R2',.01)
        original_sha = digest(self.original/'master_proposal.xlsx')
        final, budgets = consolidate(self.original, reviewed, self.base/'valid-final','DEMO-only')
        self.assertEqual(digest(self.original/'master_proposal.xlsx'), original_sha)
        persisted=json.loads((self.original/'manifest.json').read_text(encoding='utf-8'),parse_float=D)
        persisted_by_id={e['Employee_ID']:e for es in persisted['proposals'].values() for e in es}
        for e in final:
            source = persisted_by_id[e['Employee_ID']]
            for key, value in source.items():
                self.assertEqual(e[key], value)
        alpha = next(b for b in budgets if b[0]=='Team Alpha')
        self.assertEqual(alpha[4],14000)
        self.assertEqual(alpha[5],11534000)
        self.assertEqual(alpha[6],-14000)
        self.assertEqual(alpha[8],'EXCEEDED')
        self.assertEqual(sum(e['Human_Modified'] for e in final),1)
        self.assertTrue(all(e['Final_Salary'] % 100 == 0 for e in final))

    def test_unauthorized_change_rejected_and_logged(self):
        reviewed = self.reviewed_copy('tampered')
        change_cell(reviewed/'team_L-A.xlsx',2,'H2',999999)
        out = self.base/'tampered-final'
        with self.assertRaisesRegex(ValueError,'rejected'):
            consolidate(self.original,reviewed,out,'DEMO-only')
        log = json.loads((out/'validation_exceptions.json').read_text())
        self.assertTrue(any(e['error']=='PROTECTED_FIELD_CHANGED' for e in log['errors']))
        self.assertFalse((out/'consolidated_final.xlsx').exists())

    def test_formula_tamper_rejected(self):
        reviewed = self.reviewed_copy('formula-tampered')
        change_cell(reviewed/'team_L-A.xlsx',2,'T2',999999)
        with self.assertRaisesRegex(ValueError,'rejected'):
            consolidate(self.original,reviewed,self.base/'formula-final','DEMO-only')

    def test_missing_review_rejected(self):
        reviewed = self.reviewed_copy('missing')
        (reviewed/'leadership_review.xlsx').unlink()
        with self.assertRaisesRegex(ValueError,'rejected'):
            consolidate(self.original,reviewed,self.base/'missing-final','DEMO-only')

    def test_pasted_invalid_adjustment_rejected(self):
        for index, (filename, value) in enumerate([('team_L-C.xlsx',-.01),('team_L-A.xlsx',.06)]):
            reviewed = self.reviewed_copy(f'invalid-{index}')
            change_cell(reviewed/filename,2,'R2',value)
            with self.assertRaisesRegex(ValueError,'rejected'):
                consolidate(self.original,reviewed,self.base/f'invalid-final-{index}','DEMO-only')

    def test_zero_adjustment_retains_synthetic_salaries(self):
        for e in self.records:
            self.assertEqual(human_result(e,0,-.05,.05)['Final_Salary'],e['New_Salary'])

    def test_excel_recalculation_after_edit(self):
        target=self.base/'excel-check.xlsx'
        subprocess.run([os.environ.get('SALARY_NODE','node'),str(ROOT/'tests/recalculate.mjs'),str(self.original/'team_L-A.xlsx'),str(target)],check=True)
        import openpyxl
        w=openpyxl.load_workbook(target,data_only=True)
        self.assertEqual(w['Summary']['B8'].value,11534000)
        self.assertEqual(w['Summary']['B9'].value,-14000)
        self.assertEqual(w['Summary']['B11'].value,'EXCEEDED')
        self.assertEqual(w['Detail']['T2'].value,1937600)
        w.close()


if __name__=='__main__':
    unittest.main()
